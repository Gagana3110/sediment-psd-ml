"""Build the review presentation (report/slides.pptx) from the results on disk.

Usage:  python report/make_slides.py
"""
import json
import sys
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
import config as cfg

INK, MUTED, BLUE, ORANGE = (RGBColor.from_string(h) for h in ("0B0B0B", "52514E", "1C5CAB", "EB6834"))
SURFACE, RULE = RGBColor.from_string("FCFCFB"), RGBColor.from_string("E1E0D9")
W, H = Inches(13.333), Inches(7.5)
FONT = "Segoe UI"


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H

    def slide(self, title=None, kicker=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = SURFACE
        if title:
            if kicker:
                self.text(s, kicker.upper(), 0.6, 0.38, 12, 0.3, size=11, color=ORANGE, bold=True)
            self.text(s, title, 0.6, 0.65, 12.1, 0.9, size=27, bold=True)
            n = len(self.prs.slides)
            self.text(s, f"{n}", 12.3, 6.95, 0.5, 0.3, size=10, color=MUTED, align="r")
        return s

    def text(self, s, content, x, y, w, h, size=16, color=INK, bold=False, align="l"):
        """content: a string or a list of paragraphs; (text, level) tuples make bullets."""
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        items = content if isinstance(content, list) else [content]
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            level = 0
            if isinstance(item, tuple):
                item, level = item
            lead, _, rest = item.partition("||")       # "Bold lead||rest of sentence"
            if not rest:
                lead, rest = "", item
            prefix = "•  " if isinstance(content, list) else ""
            runs = [(prefix + lead, True)] if lead else []
            runs.append(((" " if lead else prefix) + rest, bold))
            for txt, b in runs:
                r = p.add_run()
                r.text = txt
                r.font.size, r.font.bold, r.font.name = Pt(size - 2 * level), b, FONT
                r.font.color.rgb = MUTED if level else color
            p.space_after = Pt(size * 0.55)
            if level:
                p.level = 1
            if align == "r":
                from pptx.enum.text import PP_ALIGN
                p.alignment = PP_ALIGN.RIGHT
        return box

    def image(self, s, name, x, y, max_w, max_h):
        path = cfg.FIG_DIR / name
        with Image.open(path) as im:
            ratio = im.width / im.height
        w, h = (max_w, max_w / ratio) if max_w / ratio <= max_h else (max_h * ratio, max_h)
        s.shapes.add_picture(str(path), Inches(x + (max_w - w) / 2), Inches(y), Inches(w), Inches(h))

    def table(self, s, rows, x, y, w, col_widths, size=13):
        shape = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                   Inches(0.42 * len(rows)))
        t = shape.table
        for j, cw in enumerate(col_widths):
            t.columns[j].width = Inches(cw)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = BLUE if i == 0 else SURFACE
                p = cell.text_frame.paragraphs[0]
                r = p.add_run()
                r.text = str(val)
                r.font.size, r.font.name, r.font.bold = Pt(size), FONT, i == 0
                r.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else INK


def main():
    s = json.loads((cfg.RESULTS_DIR / "summary.json").read_text())
    imp = pd.read_csv(cfg.RESULTS_DIR / "feature_importance.csv", index_col=0)
    protocols = list(s["r2"])
    r = lambda p, m, t: f"{s['r2'][protocols[p]][m][t]:.2f}"
    rf, svr = "Random Forest", "SVR (RBF)"
    d = Deck()

    # 1 title
    sl = d.slide()
    d.text(sl, "UE24CS352A MACHINE LEARNING  ·  MINI-PROJECT 62", 0.9, 1.9, 11.5, 0.4, size=13,
           color=ORANGE, bold=True)
    d.text(sl, "Machine learning for predicting sediment particle size distributions", 0.9, 2.4, 11.2, 2,
           size=40, bold=True)
    d.text(sl, "Team B-16   ·   Bandugula Sai Rithesh Reddy (PES1UG24AM064)   ·   Gagana C (PES1UG24AM099)",
           0.9, 4.9, 11.5, 0.4, size=16, color=MUTED)
    d.text(sl, "Reference: G. Egan, Stanford CS229, 2019   ·   Data: U.S. Geological Survey, doi:10.5066/P99Q4CHM",
           0.9, 5.4, 11.5, 0.4, size=13, color=MUTED)

    # 2 problem
    sl = d.slide("Floc size controls how fast mud settles, and nobody can predict it well", "The problem")
    d.text(sl, [
        "What happens||Suspended mud particles stick together into flocs and break apart again within "
        "seconds to minutes.",
        "Why it matters||Particle size sets settling speed, the most uncertain input to models of "
        "erosion, contaminant transport and carbon burial.",
        "The gap||Physical flocculation models are tuned to each site; none predicts particle size in general.",
        "Our task||Regression: predict two summaries of the particle size distribution from routinely "
        "measured water properties.",
        ("d50: the median particle diameter (µm)", 1),
        ("σ: the spread of the distribution (µm), the square root of the paper's variance", 1),
        "Success measure||R² between predicted and measured values on held-out data.",
    ], 0.6, 1.75, 6.6, 5, size=16)
    d.image(sl, "fig1_example_psd.png", 7.5, 2.0, 5.3, 4.2)

    # 3 paper
    sl = d.slide("The reference paper reached R² of 0.75 to 0.81, and doubted its own generality", "Reference paper")
    d.text(sl, [
        "Data||1,648 hourly samples from one instrument platform in South San Francisco Bay (2018–19). "
        "Not publicly available.",
        "Features||Salinity, wave velocity, tidal velocity, temperature, chlorophyll, refractive index "
        "and two light-absorption ratios.",
        "Method||Random Forest importance to pick the top four features, out-of-bag score to pick 32 trees, "
        "then SVR with a sweep over C and ε.",
        "Result||Random Forest R² = 0.81 (d50) and 0.75 (σ²); SVR 0.76 and 0.76. Salinity dominated.",
        "Its open question||“They could each be overfitting … rendering them less generalizable.” "
        "The models were only tested on a random split of one site's time series.",
        "Our plan||Reproduce the method on public data, then test exactly that open question.",
    ], 0.6, 1.75, 12, 5, size=17)

    # 4 dataset
    sl = d.slide("We used USGS casts from the same bay: two sites, four days, 193 files", "Dataset")
    d.text(sl, [
        "Source||USGS data release, South San Francisco Bay, summer 2020 (doi:10.5066/P99Q4CHM). "
        "Same bay and instrument family as the paper.",
        "Design||Vertical casts at a channel site (11–14 m deep) and a shallows site (2–5 m), "
        "about one minute at each depth, on 21, 22, 24 and 28 July 2020.",
        "Contents||LISST: 36-class particle size spectrum, temperature, beam attenuation. "
        "CTD: salinity, turbidity, depth.",
        "Compared with the paper||We have salinity and temperature; no chlorophyll, refractive index "
        "or wave and tidal velocity.",
    ], 0.6, 1.75, 12, 2.6, size=16)
    d.image(sl, "fig2_data_overview.png", 0.6, 4.05, 12.1, 3.1)

    # 5 pipeline
    sl = d.slide(f"Cleaning turned 45,639 raw samples into {s['n_samples']:,} reliable bursts", "Data preparation")
    d.table(sl, [
        ["Step", "What it removes or does", "Rows left"],
        ["Read 193 NetCDF files", "Decode the EPIC time format; drop one LISST file that duplicates another", "45,639"],
        ["Join CTD to LISST", "Correct the 3–7 s clock offset each day, then match second by second", "45,639"],
        ["Quality control", "In-air samples, optical transmission outside 30–98 %, readings with no CTD match", "33,108"],
        ["Burst averaging", "Average 10 s at a steady depth so one row is one stable measurement",
         f"{s['n_samples']:,}"],
        ["Targets", "Recompute d50 and σ from each averaged spectrum (r = 0.98 against the file's d50)",
         f"{s['n_samples']:,}"],
    ], 0.6, 1.8, 12.1, [2.9, 7.6, 1.6], size=13)
    d.text(sl, [
        f"Final table||{s['n_samples']:,} bursts from {s['n_casts']} casts (50 channel, 42 shallows).",
        "Features (9)||Sensor depth, height above bed, water depth (tide stage), tide rate, temperature, "
        "salinity, turbidity, beam attenuation, site.",
        "Left out on purpose||Total volume concentration, because it is computed from the size spectrum "
        "itself and would leak the answer.",
    ], 0.6, 4.75, 12.1, 2.3, size=16)

    # 6 method
    sl = d.slide("Five models, four test protocols of increasing difficulty", "Method")
    d.text(sl, [
        "Models",
        ("Random Forest and SVR, as in the paper", 1),
        ("Ridge regression as a linear baseline", 1),
        ("Gradient boosting", 1),
        ("Neural net with 10 hidden units (Sahin et al., 2017)", 1),
        "Tuning, following the paper",
        (f"Forest size from the out-of-bag curve: {s['n_trees']} trees", 1),
        (f"SVR sweep on validation casts: C = {s['svr_params']['C']:g}, ε = {s['svr_params']['epsilon']:g}", 1),
    ], 0.6, 1.75, 5.6, 5, size=18, bold=True)
    d.table(sl, [
        ["Test protocol", "What the model has not seen"],
        ["1. Random 70/30 (paper)", "Individual bursts only; their neighbours are in training"],
        ["2. Unseen casts (5-fold)", "Whole casts"],
        ["3. Unseen day", "A whole day of sampling"],
        ["4. Unseen site", "The other site entirely"],
    ], 6.5, 1.9, 6.2, [2.5, 3.7], size=13)
    d.text(sl, "Bursts a few seconds apart are nearly identical, so a random split lets the model look up "
               "the answer from a neighbour. Protocols 2 to 4 remove that shortcut.",
           6.5, 4.5, 6.2, 1.6, size=15, color=MUTED)

    # 7 importance
    top = imp[["impurity_d50_um", "impurity_sigma_um"]].mean(axis=1).sort_values(ascending=False)
    share = lambda f: f"{100 * imp.loc[f, 'impurity_d50_um']:.0f} % / {100 * imp.loc[f, 'impurity_sigma_um']:.0f} %"
    sl = d.slide("Depth, tide stage and attenuation carry the signal; salinity matters far less than in the paper",
                 "Feature importance")
    d.image(sl, "fig3_feature_importance.png", 0.6, 1.8, 12.1, 3.9)
    d.text(sl, [
        f"Sensor depth leads for d50 and water depth for σ. Salinity scores {share('salinity_psu')} "
        f"(d50 / σ) against 72 % / 56 % in the paper: over our four days it only varied by about 2 psu.",
        "Permutation importance on unseen casts rates beam attenuation much higher than the paper's method does.",
    ], 0.6, 5.85, 12.1, 1.2, size=15)

    # 8 tuning
    sl = d.slide(f"Tuning: the out-of-bag score levels off by {s['n_trees']} trees; SVR prefers a small C", "Tuning")
    d.image(sl, "fig4_oob_trees.png", 0.5, 1.8, 5.2, 3.6)
    d.image(sl, "fig5_svr_tuning.png", 5.8, 1.8, 7.1, 3.6)
    d.text(sl, [
        f"We keep the smallest forest within 0.01 of the best out-of-bag R² ({s['n_trees']} trees; the paper chose 32).",
        "SVR is tuned on casts it never saw, which favours a smooth fit (small C). "
        "The paper's very large C (2048) fits a random split better but memorises more.",
    ], 0.6, 5.6, 12.1, 1.4, size=15)

    # 9 replication
    sl = d.slide(f"With the paper's random split we match its result: R² {r(0, rf, 'd50_um')} and "
                 f"{r(0, rf, 'sigma_um')}", "Result 1: replication")
    d.image(sl, "fig6_pred_vs_true_random.png", 0.5, 1.7, 6.3, 5.4)
    d.table(sl, [
        ["Random 70/30 split", "d50", "σ"],
        ["Paper: Random Forest", "0.81", "0.75"],
        ["Ours: Random Forest", r(0, rf, "d50_um"), r(0, rf, "sigma_um")],
        ["Paper: SVR", "0.76", "0.76"],
        ["Ours: SVR", r(0, svr, "d50_um"), r(0, svr, "sigma_um")],
    ], 7.2, 1.9, 5.5, [3.1, 1.2, 1.2], size=14)
    d.text(sl, [
        "Different data and features, same headline accuracy for Random Forest.",
        "Our SVR is lower here because we tuned it for unseen casts, not for this split.",
        "Paper values for σ are for the variance σ².",
    ], 7.2, 4.4, 5.5, 2.6, size=15)

    # 10 generalisation
    sl = d.slide("On data the model has not seen, accuracy drops by a third, then collapses across sites",
                 "Result 2: generalisation")
    d.image(sl, "fig8_protocol_comparison.png", 0.4, 1.7, 8.3, 5.5)
    d.text(sl, [
        f"Unseen casts||Random Forest drops from {r(0, rf, 'd50_um')} to {r(1, rf, 'd50_um')} for d50 and "
        f"from {r(0, rf, 'sigma_um')} to {r(1, rf, 'sigma_um')} for σ.",
        f"Best honest model||SVR: {r(1, svr, 'd50_um')} / {r(1, svr, 'sigma_um')} on unseen casts, "
        f"{r(2, svr, 'd50_um')} / {r(2, svr, 'sigma_um')} on an unseen day.",
        "Ranking flips||The flexible tree models win the random split; the smoother SVR wins once "
        "neighbours are removed.",
        "Unseen site||Every model is worse than predicting the mean. The two sites barely overlap in "
        "water depth or particle size.",
    ], 8.9, 1.8, 4.0, 5.3, size=14)

    # 11 conclusions
    sl = d.slide("Conclusions: about 60 % of the variation is predictable; the split decides the score", "Conclusions")
    d.text(sl, [
        "What works||Depth, tide stage, salinity, temperature and two optical measurements explain about "
        "56–62 % of the variation in d50 and σ on new casts.",
        "What the paper's score hid||Its R² of 0.75 to 0.81 came from a random split of autocorrelated data. "
        "We reproduce it, then show it falls under grouped testing. Its overfitting concern was justified.",
        "What does not work||Transfer to a different site. A general model needs training data from many sites.",
        "Limitations||Three to keep in mind.",
        ("Four days in one summer and two sites", 1),
        ("Salinity varied by only about 2 psu, so its effect is barely tested", 1),
        ("Beam attenuation comes from the same instrument as the targets", 1),
        "Next steps||Add current and wave measurements from the release's hydrodynamic time series; "
        "train across more sites and seasons.",
    ], 0.6, 1.75, 12.1, 5.3, size=16)

    # 12 demo
    sl = d.slide("Live demo and repository", "Demo")
    d.text(sl, [
        "Rebuild everything from the raw files",
        ("python src/build_dataset.py", 1),
        ("python src/train_evaluate.py", 1),
        "Hold out one cast, retrain, compare predicted and measured",
        ("python demo.py", 1),
        ("python demo.py --cast C21", 1),
        "Predict for conditions we type in",
        ("python demo.py --predict --depth 6 --water-depth 12.5 --temp 21.5 --attn 20 --site channel", 1),
    ], 0.6, 1.75, 7.4, 5, size=17, bold=True)
    d.table(sl, [
        ["File", "Role"],
        ["src/build_dataset.py", "NetCDF → clean table"],
        ["src/psd.py", "d50 and σ from a spectrum"],
        ["src/train_evaluate.py", "Models, tuning, protocols"],
        ["src/plots.py", "Figures"],
        ["demo.py", "Live demo"],
        ["report/", "Write-up and slides"],
    ], 8.3, 1.9, 4.5, [2.3, 2.2], size=12)

    out = Path(__file__).with_name("slides.pptx")
    d.prs.save(out)
    print(f"wrote {out.relative_to(cfg.ROOT)} ({len(d.prs.slides)} slides)")


if __name__ == "__main__":
    main()
